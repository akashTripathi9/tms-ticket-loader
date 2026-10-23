"""Builds the sample 'NUPI Ticket Import Template.xlsx'.

The live page generates the same workbook with real users/projects/statuses read
from TMS after sign-in. This offline sample uses clearly marked SAMPLE names so the
layout and dropdown behaviour can be reviewed before anyone signs in.
"""
from datetime import date, timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "template" / "NUPI Ticket Import Template.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)

ACCENT = "0E7C6B"
FONT = "Arial"

COLUMNS = [
    # header, key, width, required
    ("Ticket title *", "name", 46, True),
    ("Project *", "project", 24, True),
    ("Assign to", "assignee", 24, False),
    ("Status", "status", 20, False),
    ("Priority", "priority", 14, False),
    ("Start date", "dateStart", 14, False),
    ("Due date", "dateEnd", 14, False),
    ("Estimated hours", "estimatedHours", 16, False),
    ("Teams", "teams", 22, False),
    ("Description", "description", 64, False),
]
TEAMS = ["SAMPLE – Mobile", "SAMPLE – Backend", "SAMPLE – QA"]
TODAY = date.today()
DATES = [TODAY + timedelta(days=i) for i in range(-14, 366)]  # two weeks back to a year ahead
USERS = ["SAMPLE – Akash Tripathi", "SAMPLE – Vipul Sharma", "SAMPLE – Riya Mehta", "SAMPLE – Ankit Dubey"]
PROJECTS = ["SAMPLE – NUPI App", "SAMPLE – Backoffice", "SAMPLE – Merchant Portal"]
STATUSES = ["To Do", "In Progress", "In Review", "Integration Done", "Done"]
PRIORITIES = ["Low", "Normal", "High", "Urgent"]

wb = Workbook()
ws = wb.active
ws.title = "Tickets"
lists = wb.create_sheet("Lists")
info = wb.create_sheet("Instructions")

# ---- Lists sheet -------------------------------------------------------------
for ci, (title, values) in enumerate([("Users", USERS), ("Projects", PROJECTS), ("Statuses", STATUSES), ("Priorities", PRIORITIES), ("Teams", TEAMS), ("Dates", DATES)], start=1):
    c = lists.cell(row=1, column=ci, value=title)
    c.font = Font(name=FONT, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=ACCENT)
    for r, v in enumerate(values, start=2):
        cell = lists.cell(row=r, column=ci, value=v)
        cell.font = Font(name=FONT)
        if isinstance(v, date):
            cell.number_format = "yyyy-mm-dd"
    lists.column_dimensions[get_column_letter(ci)].width = 30 if title != "Dates" else 14
lists.cell(row=8, column=8, value="These lists are SAMPLE values. The template downloaded from NUPI Ticket Loader after sign-in carries the real users, projects and statuses from TMS.").font = Font(name=FONT, italic=True, color="6B7A72")
lists.freeze_panes = "A2"

def list_ref(col_idx: int, n: int) -> str:
    L = get_column_letter(col_idx)
    return f"Lists!${L}$2:${L}${max(2, n + 1)}"

# ---- Tickets sheet -----------------------------------------------------------
thin = Side(style="thin", color="D8DED6")
for ci, (header, key, width, required) in enumerate(COLUMNS, start=1):
    c = ws.cell(row=1, column=ci, value=header)
    c.font = Font(name=FONT, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=ACCENT if required else "3A9B8B")
    c.alignment = Alignment(vertical="center")
    ws.column_dimensions[get_column_letter(ci)].width = width
ws.row_dimensions[1].height = 24
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

MAXR = 1000
def dv_list(col_key: str, formula: str, allow_blank: bool, title: str, msg: str) -> None:
    ci = next(i for i, c in enumerate(COLUMNS, start=1) if c[1] == col_key)
    L = get_column_letter(ci)
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank, showErrorMessage=True, errorTitle=title, error=msg)
    dv.add(f"{L}2:{L}{MAXR}")
    ws.add_data_validation(dv)

dv_list("project", list_ref(2, len(PROJECTS)), False, "Unknown project", "Pick a project from the dropdown.")
dv_list("assignee", list_ref(1, len(USERS)), True, "Unknown user", "Pick a person from the dropdown, or leave blank for unassigned.")
dv_list("status", list_ref(3, len(STATUSES)), True, "Unknown status", "Pick a status, or leave blank for the project's default column.")
dv_list("priority", list_ref(4, len(PRIORITIES)), True, "Unknown priority", "Pick a priority from the dropdown.")
# Teams: dropdown as a convenience, several names may be typed with commas, so no error on free text
ci = next(i for i, c in enumerate(COLUMNS, start=1) if c[1] == "teams")
L = get_column_letter(ci)
dv = DataValidation(type="list", formula1=list_ref(5, len(TEAMS)), allow_blank=True, showErrorMessage=False, showInputMessage=True, promptTitle="Teams", prompt="Pick a team from the arrow; for several, type names separated by commas.")
dv.add(f"{L}2:{L}{MAXR}")
ws.add_data_validation(dv)
# Dates: a plain .xlsx cannot show a calendar control, so each date cell gets a click-to-pick list of dates
# (two weeks back to a year ahead). Typing a date still works because the list does not reject other values.
for key in ("dateStart", "dateEnd"):
    ci = next(i for i, c in enumerate(COLUMNS, start=1) if c[1] == key)
    L = get_column_letter(ci)
    dv = DataValidation(type="list", formula1=list_ref(6, len(DATES)), allow_blank=True, showErrorMessage=False, showInputMessage=True,
                        promptTitle="Pick a date", prompt="Click the arrow and pick the date, or type one (2026-10-23 or 23/10/2026). In Google Sheets, double-click the cell for a calendar.")
    dv.add(f"{L}2:{L}{MAXR}")
    ws.add_data_validation(dv)
    for r in range(2, MAXR + 1):
        ws.cell(row=r, column=ci).number_format = "yyyy-mm-dd"
ci = next(i for i, c in enumerate(COLUMNS, start=1) if c[1] == "estimatedHours")
dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, error="Hours must be a number, e.g. 4 or 2.5.")
dv.add(f"{get_column_letter(ci)}2:{get_column_letter(ci)}{MAXR}")
ws.add_data_validation(dv)

EXAMPLES = [
    ["EXAMPLE – Add OTP resend cooldown on login screen", PROJECTS[0], USERS[1], "To Do", "High", DATES[14], DATES[18], 6, TEAMS[0],
     "Show a 30 second countdown after the first OTP, disable the resend button meanwhile, and log the resend event to analytics."],
    ["EXAMPLE – Fix settlement report export timeout", PROJECTS[1], "", "", "Normal", None, None, 8, "",
     "Export of more than 5000 rows times out. Move the export to a background job and email the file when ready."],
]
for r, row in enumerate(EXAMPLES, start=2):
    for ci, v in enumerate(row, start=1):
        c = ws.cell(row=r, column=ci, value=v)
        if isinstance(v, date):
            c.number_format = "yyyy-mm-dd"
        c.font = Font(name=FONT, italic=True, color="6B7A72")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(bottom=thin)
for r in range(4, 60):
    for ci in range(1, len(COLUMNS) + 1):
        c = ws.cell(row=r, column=ci)
        c.font = Font(name=FONT)
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws["A1"].comment = Comment("Required. A short, specific title. Rows without a title are skipped.", "NUPI Ticket Loader")
ws["B1"].comment = Comment("Required. Pick from the dropdown (list comes from TMS).", "NUPI Ticket Loader")
ws["C1"].comment = Comment("Optional. Pick a person, or leave blank to create the ticket unassigned.", "NUPI Ticket Loader")
ws["D1"].comment = Comment("Optional. Leave blank to use the project's default board column.", "NUPI Ticket Loader")
ws["E1"].comment = Comment("Optional. Low, Normal, High or Urgent.", "NUPI Ticket Loader")
ws["F1"].comment = Comment("Optional. Click the arrow to pick a date, or type 2026-10-23 / 23/10/2026.", "NUPI Ticket Loader")
ws["G1"].comment = Comment("Optional. Click the arrow to pick a date, or type 2026-10-23 / 23/10/2026. Must not be before the start date.", "NUPI Ticket Loader")
ws["H1"].comment = Comment("Optional. Number of hours, e.g. 4 or 2.5.", "NUPI Ticket Loader")
ws["I1"].comment = Comment("Optional. One team from the dropdown, or several names separated by commas.", "NUPI Ticket Loader")
ws["J1"].comment = Comment("Recommended. What needs to be done and how we will know it is done.", "NUPI Ticket Loader")

# ---- Instructions sheet --------------------------------------------------------
info.column_dimensions["A"].width = 110
LINES = [
    ("How to use this sheet", 14, True),
    ("1. Fill one ticket per row on the “Tickets” sheet. Delete the two grey EXAMPLE rows before importing.", 11, False),
    ("2. Columns marked * are required: Ticket title and Project. Everything else is optional.", 11, False),
    ("3. Project, Assign to, Status, Priority and Teams are dropdowns fed by the “Lists” sheet. In this sample the lists are placeholders; the template downloaded from NUPI Ticket Loader after sign-in carries the real names from TMS and every extra field that exists on the TMS ticket form (e.g. ticket type, reviewer).", 11, False),
    ("4. Start date / Due date: click the cell, then the arrow, and pick the date from the list (two weeks back to a year ahead). Typing works too: 2026-10-23 or 23/10/2026. In Google Sheets, double-click a date cell for a calendar.", 11, False),
    ("5. Leave Status blank to use the project's default column. Leave Assign to blank to create the ticket unassigned; you can assign later in TMS. Estimated hours is a plain number (4, 2.5).", 11, False),
    ("6. In NUPI Ticket Loader, click “Import Excel / CSV” and choose this file. Check the preview, fix anything marked red, click Review, then Confirm. Nothing is created in TMS until you confirm.", 11, False),
    ("", 11, False),
    ("Legend", 12, True),
    ("Dark green header = required column. Light green header = optional. Grey italic rows = examples to delete.", 11, False),
]
for r, (text, size, bold) in enumerate(LINES, start=1):
    c = info.cell(row=r, column=1, value=text)
    c.font = Font(name=FONT, size=size, bold=bold)
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print(OUT)
